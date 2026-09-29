"""
export_engine.py  —  Batch export: Excel pipeline tracker + CSV + self-contained HTML report
"""

from __future__ import annotations
import io
import csv
import json
import base64
from datetime import datetime
from typing import Any


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL EXPORT
# ══════════════════════════════════════════════════════════════════════════════
def generate_excel_export(results: list[dict]) -> io.BytesIO:
    """
    Generate a professional Excel workbook with:
      Sheet 1: Ranked Leaderboard (one row per candidate, colour-coded)
      Sheet 2: Skill Gap Matrix (candidates × skills heatmap)
      Sheet 3: AI Recommendations (all suggestions expanded)
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side, GradientFill
        )
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import ColorScaleRule
    except ImportError:
        raise ImportError("openpyxl required: pip install openpyxl")

    wb = Workbook()

    # ── Colour palette ──────────────────────────────────────────────────────
    C_HEADER_BG   = "0D1117"
    C_HEADER_FG   = "00D4FF"
    C_GREEN_BG    = "0D2B1F"
    C_GREEN_FG    = "10B981"
    C_AMBER_BG    = "2B1F0D"
    C_AMBER_FG    = "F59E0B"
    C_RED_BG      = "2B0D0D"
    C_RED_FG      = "EF4444"
    C_ALT_ROW     = "0F1724"
    C_WHITE       = "E2E8F0"
    C_MUTED       = "64748B"
    C_BORDER      = "1E2D45"

    def header_cell(ws, row, col, value, width=18):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True, color=C_HEADER_FG, name="Arial", size=10)
        cell.fill = PatternFill("solid", fgColor=C_HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=C_BORDER))
        ws.column_dimensions[get_column_letter(col)].width = width
        return cell

    def tier_colors(tier_code: str):
        return {
            "green": (C_GREEN_BG, C_GREEN_FG),
            "amber": (C_AMBER_BG, C_AMBER_FG),
            "red":   (C_RED_BG, C_RED_FG),
        }.get(tier_code, (C_ALT_ROW, C_WHITE))

    ranked = sorted(results, key=lambda r: r.get("score", 0), reverse=True)

    # ════════════════════════════════════════════════════════════════════════
    # SHEET 1: Leaderboard
    # ════════════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "Candidate Leaderboard"
    ws1.sheet_view.showGridLines = False
    ws1.row_dimensions[1].height = 40

    headers = [
        ("Rank", 6), ("Candidate", 28), ("Match Score", 13), ("Quality", 10),
        ("Fit Verdict", 20), ("Experience", 12), ("Missing Skills", 35),
        ("Certifications", 22), ("Strengths", 40), ("Key Concern", 40),
    ]
    for col, (h, w) in enumerate(headers, 1):
        header_cell(ws1, 1, col, h, w)

    for row_idx, r in enumerate(ranked, 2):
        rank = row_idx - 1
        verdict = r.get("fit_verdict", {})
        tier_code = verdict.get("tier_code", "amber")
        bg, fg = tier_colors(tier_code)

        row_fill = PatternFill("solid", fgColor=C_ALT_ROW if rank % 2 == 0 else C_HEADER_BG)

        def cell(col, value, bold=False, color=C_WHITE, fill=None, wrap=False):
            c = ws1.cell(row=row_idx, column=col, value=value)
            c.font = Font(name="Arial", size=10, bold=bold, color=color)
            c.fill = fill or row_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
            return c

        cell(1, rank, bold=True)
        cell(2, r.get("name", "").replace(".pdf", ""), bold=True).alignment = Alignment(horizontal="left", vertical="center")
        
        score = r.get("score", 0)
        score_color = C_GREEN_FG if score >= 70 else (C_AMBER_FG if score >= 45 else C_RED_FG)
        cell(3, f"{score}%", bold=True, color=score_color)
        cell(4, f"{r.get('quality', 0)}/10")

        verdict_cell = cell(5, verdict.get("tier", "—"), bold=True, color=fg,
                            fill=PatternFill("solid", fgColor=bg))

        cell(6, f"{r.get('experience_years', '?')} yrs")
        
        missing_str = ", ".join(r.get("missing", [])[:6])
        c7 = ws1.cell(row=row_idx, column=7, value=missing_str or "None")
        c7.font = Font(name="Arial", size=9, color=C_RED_FG if missing_str else C_GREEN_FG)
        c7.fill = row_fill
        c7.alignment = Alignment(wrap_text=True, vertical="center")

        certs = ", ".join(r.get("certifications", [])) or "None"
        cell(8, certs)

        strengths = verdict.get("strengths", [])
        cell(9, " • ".join(strengths[:2]) if strengths else "—", wrap=True).alignment = Alignment(horizontal="left", wrap_text=True, vertical="center")

        concerns = verdict.get("concerns", [])
        cell(10, concerns[0] if concerns else "—", wrap=True).alignment = Alignment(horizontal="left", wrap_text=True, vertical="center")

        ws1.row_dimensions[row_idx].height = 32

    # Freeze header
    ws1.freeze_panes = "A2"

    # ════════════════════════════════════════════════════════════════════════
    # SHEET 2: Skill Gap Matrix
    # ════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Skill Gap Matrix")
    ws2.sheet_view.showGridLines = False

    # Collect all unique skills across all candidates from analysis
    all_skills: list[str] = []
    seen = set()
    for r in ranked:
        for cat_data in r.get("analysis", {}).values():
            for s in cat_data.get("found", []) + cat_data.get("missing", []):
                if s not in seen:
                    all_skills.append(s)
                    seen.add(s)
    all_skills = all_skills[:30]  # cap at 30 skills

    # Headers: candidate names
    ws2.cell(1, 1, "Skill").font = Font(bold=True, color=C_HEADER_FG, name="Arial")
    ws2.cell(1, 1).fill = PatternFill("solid", fgColor=C_HEADER_BG)
    ws2.column_dimensions["A"].width = 24

    for col, r in enumerate(ranked, 2):
        name = r.get("name", "").replace(".pdf", "")[:18]
        c = ws2.cell(1, col, name)
        c.font = Font(bold=True, color=C_HEADER_FG, name="Arial", size=9)
        c.fill = PatternFill("solid", fgColor=C_HEADER_BG)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws2.column_dimensions[get_column_letter(col)].width = 16

    for skill_row, skill in enumerate(all_skills, 2):
        ws2.cell(skill_row, 1, skill.title()).font = Font(name="Arial", size=9, color=C_WHITE)
        ws2.cell(skill_row, 1).fill = PatternFill("solid", fgColor=C_ALT_ROW if skill_row % 2 == 0 else C_HEADER_BG)

        for col, r in enumerate(ranked, 2):
            # Check if skill is found
            found_skills = set()
            missing_skills_set = set()
            for cat_data in r.get("analysis", {}).values():
                found_skills.update(cat_data.get("found", []))
                missing_skills_set.update(cat_data.get("missing", []))

            if skill in found_skills:
                val, bg_c, fg_c = "✓", C_GREEN_BG, C_GREEN_FG
            elif skill in missing_skills_set:
                val, bg_c, fg_c = "✗", C_RED_BG, C_RED_FG
            else:
                val, bg_c, fg_c = "—", C_HEADER_BG, C_MUTED

            c = ws2.cell(skill_row, col, val)
            c.font = Font(bold=True, color=fg_c, name="Arial", size=11)
            c.fill = PatternFill("solid", fgColor=bg_c)
            c.alignment = Alignment(horizontal="center", vertical="center")

        ws2.row_dimensions[skill_row].height = 22

    ws2.freeze_panes = "B2"

    # ════════════════════════════════════════════════════════════════════════
    # SHEET 3: AI Recommendations
    # ════════════════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("AI Recommendations")
    ws3.sheet_view.showGridLines = False

    rec_headers = [("Candidate", 28), ("Priority", 12), ("Category", 18),
                   ("Finding", 55), ("Action", 55)]
    for col, (h, w) in enumerate(rec_headers, 1):
        header_cell(ws3, 1, col, h, w)

    priority_colors = {"critical": C_RED_FG, "high": C_AMBER_FG, "medium": C_HEADER_FG, "low": C_GREEN_FG}
    rec_row = 2
    for r in ranked:
        for sug in r.get("suggestions", []):
            if isinstance(sug, dict):
                priority = sug.get("priority", "medium")
                p_color = priority_colors.get(priority, C_WHITE)
                row_fill = PatternFill("solid", fgColor=C_ALT_ROW if rec_row % 2 == 0 else C_HEADER_BG)

                ws3.cell(rec_row, 1, r.get("name", "").replace(".pdf", "")).font = Font(name="Arial", size=9, color=C_WHITE, bold=True)
                ws3.cell(rec_row, 1).fill = row_fill

                pc = ws3.cell(rec_row, 2, priority.upper())
                pc.font = Font(name="Arial", size=9, bold=True, color=p_color)
                pc.fill = row_fill
                pc.alignment = Alignment(horizontal="center")

                ws3.cell(rec_row, 3, sug.get("category", "")).font = Font(name="Arial", size=9, color=C_MUTED)
                ws3.cell(rec_row, 3).fill = row_fill

                mc = ws3.cell(rec_row, 4, sug.get("message", ""))
                mc.font = Font(name="Arial", size=9, color=C_WHITE)
                mc.fill = row_fill
                mc.alignment = Alignment(wrap_text=True, vertical="top")

                ac = ws3.cell(rec_row, 5, sug.get("action", ""))
                ac.font = Font(name="Arial", size=9, color=C_MUTED, italic=True)
                ac.fill = row_fill
                ac.alignment = Alignment(wrap_text=True, vertical="top")

                ws3.row_dimensions[rec_row].height = 36
                rec_row += 1

    ws3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════════
# CSV EXPORT
# ══════════════════════════════════════════════════════════════════════════════
def generate_csv_export(results: list[dict]) -> io.BytesIO:
    """Flat CSV of all candidates — ready for ATS import or spreadsheet tools."""
    ranked = sorted(results, key=lambda r: r.get("score", 0), reverse=True)
    buf = io.StringIO()
    writer = csv.writer(buf)

    writer.writerow([
        "Rank", "Candidate", "Match Score (%)", "Quality Score (/10)",
        "Fit Verdict", "Experience (years)", "Missing Skills Count",
        "Missing Skills", "Extra Skills", "Certifications",
        "Skills Score (%)", "Experience Score (%)", "Projects Score (%)", "Education Score (%)",
        "AI Narrative (first 200 chars)",
    ])

    for rank, r in enumerate(ranked, 1):
        ss = r.get("section_scores", {})
        verdict = r.get("fit_verdict", {})
        writer.writerow([
            rank,
            r.get("name", "").replace(".pdf", ""),
            r.get("score", 0),
            r.get("quality", 0),
            verdict.get("tier", "—"),
            r.get("experience_years", "?"),
            len(r.get("missing", [])),
            "; ".join(r.get("missing", [])[:10]),
            "; ".join(r.get("extra_skills", [])[:8]),
            "; ".join(r.get("certifications", [])),
            ss.get("skills", "—"),
            ss.get("experience", "—"),
            ss.get("projects", "—"),
            ss.get("education", "—"),
            r.get("narrative_summary", "")[:200].replace("\n", " "),
        ])

    result = buf.getvalue().encode("utf-8")
    return io.BytesIO(result)


# ══════════════════════════════════════════════════════════════════════════════
# SELF-CONTAINED HTML REPORT
# ══════════════════════════════════════════════════════════════════════════════
def generate_html_report(results: list[dict], job_description: str = "") -> io.BytesIO:
    """
    Generate a self-contained, single-file HTML report that can be
    emailed or shared without any server. All CSS/JS is inline.
    """
    ranked = sorted(results, key=lambda r: r.get("score", 0), reverse=True)
    generated_at = datetime.now().strftime("%d %B %Y, %H:%M")

    TIER_STYLES = {
        "green":  ("Strongly Recommended", "#10b981", "#0d2b1f"),
        "amber":  ("Consider",             "#f59e0b", "#2b1f0d"),
        "red":    ("Reject",               "#ef4444", "#2b0d0d"),
    }

    def score_color(s):
        return "#10b981" if s >= 70 else ("#f59e0b" if s >= 45 else "#ef4444")

    def candidate_card(rank: int, r: dict) -> str:
        score = r.get("score", 0)
        quality = r.get("quality", 0)
        verdict = r.get("fit_verdict", {})
        tier_code = verdict.get("tier_code", "amber")
        tier_label, tier_color, tier_bg = TIER_STYLES.get(tier_code, TIER_STYLES["amber"])
        ss = r.get("section_scores", {})
        name = r.get("name", "").replace(".pdf", "")

        # Section bars
        section_bars = ""
        for sec in ["skills", "experience", "projects", "education"]:
            val = ss.get(sec, 0)
            col = score_color(val)
            section_bars += f"""
            <div class="sec-row">
              <span class="sec-label">{sec.title()}</span>
              <div class="bar-track">
                <div class="bar-fill" style="width:{val}%;background:{col}"></div>
              </div>
              <span class="sec-val" style="color:{col}">{val:.0f}%</span>
            </div>"""

        # Skill chips
        found_chips = "".join(
            f'<span class="chip chip-found">{s}</span>'
            for cat in r.get("analysis", {}).values()
            for s in cat.get("found", [])
        )
        missing_chips = "".join(
            f'<span class="chip chip-missing">{s}</span>'
            for cat in r.get("analysis", {}).values()
            for s in cat.get("missing", [])
        )
        extra_chips = "".join(
            f'<span class="chip chip-extra">{s}</span>'
            for s in r.get("extra_skills", [])[:8]
        )

        # Suggestions
        priority_colors = {"critical":"#ef4444","high":"#f59e0b","medium":"#00d4ff","low":"#10b981"}
        sug_html = ""
        for sug in r.get("suggestions", []):
            if isinstance(sug, dict):
                p = sug.get("priority", "medium")
                col = priority_colors.get(p, "#94a3b8")
                sug_html += f"""
                <div class="sug-item" style="border-left-color:{col}">
                  <div class="sug-pri" style="color:{col}">{p.upper()} · {sug.get('category','')}</div>
                  <div class="sug-msg">{sug.get('message','')}</div>
                  <div class="sug-act">→ {sug.get('action','')}</div>
                </div>"""

        # Narrative
        narrative = r.get("narrative_summary", "").replace("\n\n", "</p><p>")
        narrative_html = f"<p>{narrative}</p>" if narrative else ""

        # Verdict details
        strengths = verdict.get("strengths", [])
        concerns = verdict.get("concerns", [])
        str_html = "".join(f"<li>{s}</li>" for s in strengths)
        con_html = "".join(f"<li>{c}</li>" for c in concerns)

        rank_badge = {1:"🥇",2:"🥈",3:"🥉"}.get(rank, f"#{rank}")

        return f"""
<div class="cand-card" id="candidate-{rank}">
  <div class="cand-header">
    <div>
      <div class="cand-rank">{rank_badge}</div>
      <div class="cand-name">{name}</div>
      <div class="cand-meta">
        Exp: {r.get('experience_years','?')} yrs &nbsp;·&nbsp;
        Quality: {quality}/10 &nbsp;·&nbsp;
        {len(r.get('missing',[]))} skills gap
      </div>
    </div>
    <div style="text-align:right">
      <div class="score-big" style="color:{score_color(score)}">{score}%</div>
      <div class="score-sub">MATCH</div>
      <div class="verdict-badge" style="color:{tier_color};background:{tier_bg}">{tier_label}</div>
    </div>
  </div>

  <div class="two-col">
    <div>
      <div class="section-title">SECTION SCORES</div>
      {section_bars}
    </div>
    <div>
      <div class="section-title">FIT VERDICT</div>
      <p style="color:#94a3b8;font-size:0.85rem;margin:0 0 0.5rem">{verdict.get('reason','')}</p>
      {'<div class="section-title" style="margin-top:0.8rem">STRENGTHS</div><ul class="verdict-list" style="color:#10b981">' + str_html + '</ul>' if strengths else ''}
      {'<div class="section-title" style="margin-top:0.8rem">CONCERNS</div><ul class="verdict-list" style="color:#f59e0b">' + con_html + '</ul>' if concerns else ''}
    </div>
  </div>

  {'<div class="section-title">AI RECRUITER SUMMARY</div><div class="narrative">' + narrative_html + '</div>' if narrative_html else ''}

  <div class="section-title">SKILL ANALYSIS</div>
  <div style="margin-bottom:0.5rem">{found_chips}{missing_chips}</div>
  {'<div class="section-title" style="margin-top:0.5rem">HIDDEN STRENGTHS</div><div>' + extra_chips + '</div>' if extra_chips else ''}

  <div class="section-title" style="margin-top:1rem">AI RECOMMENDATIONS</div>
  {sug_html}

  <details style="margin-top:1rem">
    <summary style="cursor:pointer;color:#64748b;font-size:0.75rem;font-family:monospace">
      ATS KEYWORD DENSITY
    </summary>
    <div class="ats-grid">
      {''.join(f"""<div class="ats-row {'ats-present' if d['present'] else 'ats-missing'}">
        <span class="ats-term">{term}</span>
        <span class="ats-jd">JD:{d['jd_count']}</span>
        <span class="ats-res">CV:{d['resume_count']}</span>
      </div>""" for term, d in list(r.get('ats_density', {}).items())[:20])}
    </div>
  </details>
</div>"""

    # Build all cards
    cards_html = "\n".join(candidate_card(i + 1, r) for i, r in enumerate(ranked))

    # Summary table
    summary_rows = ""
    for i, r in enumerate(ranked, 1):
        verdict = r.get("fit_verdict", {})
        tc = verdict.get("tier_code", "amber")
        _, tcol, tbg = TIER_STYLES.get(tc, TIER_STYLES["amber"])
        sc = r.get("score", 0)
        summary_rows += f"""
        <tr>
          <td>{i}</td>
          <td><strong>{r.get('name','').replace('.pdf','')}</strong></td>
          <td style="color:{score_color(sc)};font-weight:700">{sc}%</td>
          <td>{r.get('quality',0)}/10</td>
          <td><span style="color:{tcol};background:{tbg};padding:2px 8px;border-radius:4px;font-size:0.75rem">{verdict.get('tier','—')}</span></td>
          <td>{r.get('experience_years','?')} yrs</td>
          <td style="color:#ef4444">{', '.join(r.get('missing',[])[:4]) or '—'}</td>
        </tr>"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>CORE-AI Resume Intelligence Report</title>
<style>
  *{{box-sizing:border-box;margin:0;padding:0}}
  :root{{
    --bg:#07090f;--card:#111827;--border:#1e2d45;
    --accent:#00d4ff;--text:#e2e8f0;--muted:#64748b;
  }}
  body{{background:var(--bg);color:var(--text);font-family:'Segoe UI',system-ui,sans-serif;padding:2rem 1rem}}
  h1{{font-family:monospace;color:var(--accent);letter-spacing:4px;text-align:center;font-size:1.8rem;margin-bottom:0.3rem}}
  .subtitle{{text-align:center;color:var(--muted);font-family:monospace;font-size:0.75rem;letter-spacing:2px;margin-bottom:2rem}}
  .meta{{text-align:center;color:var(--muted);font-size:0.75rem;margin-bottom:1rem}}
  .summary-table{{width:100%;border-collapse:collapse;margin-bottom:2rem;font-size:0.85rem}}
  .summary-table th{{background:var(--card);color:var(--accent);font-family:monospace;font-size:0.7rem;letter-spacing:1px;padding:0.6rem 0.8rem;border-bottom:1px solid var(--border);text-align:left}}
  .summary-table td{{padding:0.6rem 0.8rem;border-bottom:1px solid var(--border);color:var(--text)}}
  .summary-table tr:hover td{{background:#0d1220}}
  .cand-card{{background:var(--card);border:1px solid var(--border);border-radius:12px;padding:1.5rem;margin-bottom:1.5rem;border-top:3px solid var(--accent)}}
  .cand-header{{display:flex;justify-content:space-between;align-items:flex-start;margin-bottom:1.2rem;border-bottom:1px solid var(--border);padding-bottom:1rem}}
  .cand-rank{{font-size:1.5rem;margin-bottom:0.2rem}}
  .cand-name{{font-size:1.2rem;font-weight:700;color:var(--text);margin-bottom:0.2rem}}
  .cand-meta{{font-size:0.78rem;color:var(--muted)}}
  .score-big{{font-family:monospace;font-size:2.5rem;font-weight:700;line-height:1}}
  .score-sub{{font-family:monospace;font-size:0.65rem;letter-spacing:2px;color:var(--muted);margin-top:0.2rem}}
  .verdict-badge{{display:inline-block;padding:4px 12px;border-radius:20px;font-size:0.75rem;font-weight:700;margin-top:0.5rem}}
  .two-col{{display:grid;grid-template-columns:1fr 1fr;gap:1.5rem;margin-bottom:1.2rem}}
  .section-title{{font-family:monospace;font-size:0.65rem;letter-spacing:2px;color:var(--muted);margin-bottom:0.5rem;text-transform:uppercase}}
  .sec-row{{display:flex;align-items:center;gap:0.5rem;margin-bottom:0.3rem}}
  .sec-label{{font-size:0.78rem;color:var(--muted);width:80px;flex-shrink:0}}
  .bar-track{{flex:1;height:6px;background:#1e2d45;border-radius:3px;overflow:hidden}}
  .bar-fill{{height:100%;border-radius:3px;transition:width 0.5s}}
  .sec-val{{font-family:monospace;font-size:0.75rem;width:36px;text-align:right}}
  .verdict-list{{padding-left:1rem;font-size:0.82rem;margin-top:0.3rem}}
  .verdict-list li{{margin-bottom:0.2rem}}
  .narrative{{background:#0d1220;border-left:3px solid var(--accent);padding:0.8rem 1rem;border-radius:0 8px 8px 0;margin-bottom:1rem;font-size:0.87rem;line-height:1.6;color:#94a3b8}}
  .narrative p{{margin-bottom:0.6rem}}
  .chip{{display:inline-block;padding:2px 8px;border-radius:12px;font-size:0.72rem;margin:2px;font-family:monospace}}
  .chip-found{{background:rgba(16,185,129,.15);color:#34d399;border:1px solid rgba(16,185,129,.3)}}
  .chip-missing{{background:rgba(239,68,68,.15);color:#f87171;border:1px solid rgba(239,68,68,.3)}}
  .chip-extra{{background:rgba(124,58,237,.15);color:#a78bfa;border:1px solid rgba(124,58,237,.3)}}
  .sug-item{{border-left:3px solid;padding:0.5rem 0.8rem;margin:0.4rem 0;border-radius:0 6px 6px 0}}
  .sug-pri{{font-size:0.68rem;font-family:monospace;font-weight:700;margin-bottom:0.2rem}}
  .sug-msg{{font-size:0.83rem;color:#cbd5e1;margin-bottom:0.15rem}}
  .sug-act{{font-size:0.78rem;color:var(--muted);font-style:italic}}
  .ats-grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(180px,1fr));gap:4px;margin-top:0.5rem}}
  .ats-row{{display:flex;justify-content:space-between;padding:3px 6px;border-radius:4px;font-size:0.72rem}}
  .ats-present{{background:rgba(16,185,129,.1);color:#34d399}}
  .ats-missing{{background:rgba(239,68,68,.1);color:#f87171}}
  .ats-term{{font-weight:600;flex:1}}
  .ats-jd,.ats-res{{font-family:monospace;margin-left:4px;opacity:0.7}}
  @media(max-width:700px){{.two-col{{grid-template-columns:1fr}}.cand-header{{flex-direction:column;gap:1rem}}}}
  @media print{{body{{background:#fff;color:#000}}.cand-card{{break-inside:avoid}}}}
</style>
</head>
<body>
<h1>CORE-AI // RESUME INTELLIGENCE</h1>
<div class="subtitle">AI-POWERED CANDIDATE SCREENING REPORT</div>
<div class="meta">Generated {generated_at} &nbsp;·&nbsp; {len(ranked)} candidate(s) analysed</div>

<table class="summary-table">
  <thead><tr>
    <th>#</th><th>Candidate</th><th>Match</th><th>Quality</th>
    <th>Verdict</th><th>Experience</th><th>Top Missing Skills</th>
  </tr></thead>
  <tbody>{summary_rows}</tbody>
</table>

{cards_html}

<div style="text-align:center;margin-top:2rem;font-family:monospace;font-size:0.65rem;color:#1e2d45">
  CORE-AI RESUME INTELLIGENCE PLATFORM &nbsp;·&nbsp; CONFIDENTIAL
</div>
</body>
</html>"""

    return io.BytesIO(html.encode("utf-8"))
